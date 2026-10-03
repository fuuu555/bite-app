from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"D:\python\bite\BiteMap_系統目標與功能架構.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name="Noto Sans CJK TC", size=None, bold=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.25):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def add_body(doc, text, after=7):
    p = doc.add_paragraph(style="Body Text")
    p.add_run(text)
    set_paragraph_spacing(p, after=after)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    set_paragraph_spacing(p, after=3, line=1.2)
    return p


def add_numbered(doc, text, number):
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.left_indent = Inches(0.24)
    p.paragraph_format.first_line_indent = Inches(-0.24)
    p.add_run(f"{number}.  ")
    p.add_run(text)
    set_paragraph_spacing(p, after=3, line=1.2)
    return p


def add_heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    set_paragraph_spacing(p, before=12 if level == 1 else 8, after=5, line=1.1)
    return p


def style_table(table, header_fill="2F5D62"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margin(cell)
            if row_index == 0:
                set_cell_shading(cell, header_fill)
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F4F7F7")
            for paragraph in cell.paragraphs:
                set_paragraph_spacing(paragraph, after=1, line=1.15)
                for run in paragraph.runs:
                    set_run_font(run, size=9.5, bold=(row_index == 0), color=("FFFFFF" if row_index == 0 else "000000"))


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    style_table(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Noto Sans CJK TC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK TC")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    body = styles["Body Text"]
    body.font.name = "Noto Sans CJK TC"
    body._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK TC")
    body.font.size = Pt(10.5)
    body.font.color.rgb = RGBColor(0, 0, 0)

    for name, size in (("Title", 24), ("Subtitle", 13), ("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11.5)):
        style = styles[name]
        style.font.name = "Noto Sans CJK TC"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK TC")
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style_p_pr = style._element.find(qn("w:pPr"))
        if style_p_pr is not None:
            style_borders = style_p_pr.find(qn("w:pBdr"))
            if style_borders is not None:
                style_p_pr.remove(style_borders)
    styles["Title"].paragraph_format.space_after = Pt(8)
    styles["Subtitle"].paragraph_format.space_after = Pt(20)


def add_title_page(doc):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run("BiteMap 系統目標與功能架構")
    remove_paragraph_borders(p)
    p = doc.add_paragraph(style="Subtitle")
    p.add_run("美食探索 旅遊地圖 行程規劃與社交約飯服務")

    add_heading(doc, "文件說明", 1)
    add_body(doc, "BiteMap 是一套結合美食探索、旅遊地圖、行程規劃與社交約飯的旅遊服務系統。本文件整理系統的服務目標、整體架構、使用者類型、主要功能模組與資料來源，說明 BiteMap 如何將找店、規劃行程及旅遊社交整合在同一套服務中。")
    add_body(doc, "本系統的核心設計，是將觀光署官方資料、Google 評分與評論數，以及 BiteMap 自身累積的再訪率、留言與在地使用資料分開呈現。使用者可以同時參考官方資料、大眾評價與實際使用經驗，再自行判斷是否值得前往。")

    add_heading(doc, "系統重點摘要", 2)
    add_numbered(doc, "降低旅客搜尋與比較餐廳、景點及住宿所需的時間。", 1)
    add_numbered(doc, "透過多來源資料提供完整且貼近實際使用經驗的資訊。", 2)
    add_numbered(doc, "將找店、排行程及單人旅遊社交整合在同一套系統中。", 3)


def build_document():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    add_title_page(doc)

    add_heading(doc, "1 系統目標", 1)
    add_body(doc, "一般旅客在陌生地區旅遊時，往往需要分別查看官方觀光網站、Google Maps、社群評論及其他旅遊平台，才能完成餐廳、景點、住宿與評論資訊的搜尋與比較。BiteMap 希望將這些原本分散的資訊集中於同一個系統，讓使用者可以直接完成搜尋、比較、選擇及安排行程。")
    add_body(doc, "BiteMap 不只提供地點清單，也重視旅客做決策時需要的不同角度資訊。系統保留 Google 評價與 BiteMap 評價的來源差異，不將兩者直接合成一個難以理解的總分；同時將再訪率作為重要的店家判斷資訊，讓使用者了解大眾口碑與實際使用者經驗之間的差異。")

    add_heading(doc, "系統要解決的問題", 2)
    add_numbered(doc, "降低旅客搜尋與比較餐廳、景點及住宿所需的時間。", 1)
    add_numbered(doc, "透過多來源資料提供更完整且貼近實際使用經驗的資訊。", 2)
    add_numbered(doc, "將找店、排行程及單人旅遊社交整合在同一套系統中。", 3)

    add_heading(doc, "2 系統架構", 1)
    add_body(doc, "BiteMap 整體功能以使用者旅遊流程為核心，主要由探索 Explore、地圖 Map、約飯 Meal、聊天室 Chat 及個人頁面 Profile 五大模組構成。各模組彼此串接，形成從探索地點到實際執行旅遊與用餐計畫的完整流程，而不是各自獨立運作。")
    doc.add_page_break()
    add_table(doc,
              ["模組", "核心用途", "主要串接"],
              [
                  ("探索 Explore", "依距離、價格與料理分類協助使用者快速找店並比較結果。", "餐廳詳細頁 收藏 行程 約飯"),
                  ("地圖 Map", "整合 BiteMap 店家與觀光署餐飲、景點、住宿及服務站資料。", "地點篩選 我的行程 路線查看"),
                  ("約飯 Meal", "建立公開或私人飯局，讓旅客共同討論並投票選擇餐廳。", "餐廳詳細頁 約飯聊天室"),
                  ("聊天室 Chat", "支援一對一聊天與約飯群組即時溝通。", "好友 追蹤 約飯成員"),
                  ("個人頁面 Profile", "管理個人資料、收藏、留言、再訪及社交資訊。", "收藏紀錄 留言 再訪 好友"),
              ],
              widths=[1.35, 3.1, 2.0])

    add_heading(doc, "主要使用流程", 2)
    add_body(doc, "使用者登入後，可以從探索或地圖尋找餐廳、景點與住宿，接著查看 Google 評價、BiteMap 再訪率及店家資訊，再依需求收藏、加入行程或建立約飯。完成行程調整或社交互動後，使用者可以依照規劃執行旅遊與用餐計畫。")
    add_body(doc, "登入 → 探索餐廳或使用地圖尋找地點 → 查看 Google 評價、BiteMap 再訪率及店家資訊 → 收藏、加入行程或建立約飯 → 調整旅遊行程或與其他使用者互動 → 實際執行旅遊與用餐計畫")
    add_body(doc, "行程功能直接與地圖整合。使用者看到餐廳、景點或住宿後，可以立即加入自己的旅遊行程，再依需求調整順序，並於 BiteMap 地圖查看行程位置與順序線路，形成從探索到實際行程的完整流程。")

    add_heading(doc, "3 系統使用者", 1)
    add_heading(doc, "3.1 一般旅客", 2)
    add_body(doc, "一般旅客是 BiteMap 最主要的使用族群，包含到陌生城市旅遊、需要快速尋找餐廳、住宿或景點的旅客，也包含不知道要吃什麼、希望系統協助縮小選擇範圍的使用者。")
    add_bullet(doc, "希望比較 Google 大眾評價與實際再訪經驗的使用者。")
    add_bullet(doc, "想將餐廳、住宿與景點整理成完整旅遊行程的旅客。")
    add_bullet(doc, "一個人旅遊，希望尋找其他人共同用餐的使用者。")
    add_bullet(doc, "希望透過留言、收藏、追蹤及好友功能記錄與分享美食經驗的使用者。")
    add_body(doc, "BiteMap 特別考量單人旅遊情境。一般美食地圖通常只能解決去哪裡吃的問題，而 BiteMap 的約飯功能進一步讓使用者能在選定餐廳後建立飯局，尋找其他旅客一起用餐。")

    add_heading(doc, "3.2 系統管理員", 2)
    add_body(doc, "管理員主要負責系統資料品質與平台內容管理，包括 BiteMap 店家新增與維護、料理分類管理、觀光署資料匯入與管理，以及店家發布狀態管理。管理員專注於資料維護與平台管理，一般使用者則專注於搜尋、旅遊規劃與社交互動。")

    add_heading(doc, "4 主要功能模組", 1)
    add_heading(doc, "4.1 多來源資料整合", 2)
    add_body(doc, "BiteMap 不只依賴單一來源判斷一家店是否值得前往。系統將觀光署官方資料作為餐飲、景點、住宿及旅遊服務站等旅遊據點的重要來源，再搭配 Google 評分與評論數，以及 BiteMap 自身累積的使用者留言、再訪率與實際使用紀錄。")
    add_body(doc, "例如，同一家餐廳可以同時呈現 Google 4.7 分、2,800 則評論，以及 BiteMap 72% 使用者願意再訪。這兩類資料不會被混合成一個不透明的分數，而是保留各自的來源與數據，使旅客可以自行比較。")
    add_body(doc, "官方資料提供完整、可信的旅遊據點資訊；Google 提供大量大眾評價；BiteMap 的再訪率與留言則反映實際使用者吃過之後是否還願意再次前往。透過三種不同角度的資訊交叉參考，資料更接近旅客真正做決策時需要的內容。")

    add_heading(doc, "4.2 探索 Explore", 2)
    add_body(doc, "探索功能主要解決不知道要吃什麼，以及資料太多、不知道如何選擇的問題。使用者可以依照距離、價格及料理分類設定條件，系統從符合條件的餐廳中進行排序，並優先顯示 Top 3。Top 3 下方仍保留完整結果，因此使用者可以快速選擇，也可以繼續比較其他店家。")
    add_body(doc, "餐廳卡片同時呈現 Google 評分、Google 評論數、BiteMap 再訪率、價格、料理類型與距離等資訊，降低使用者在多個應用程式之間反覆切換搜尋的需求。探索功能的核心價值，是將搜尋、篩選、比較與決策集中在同一個流程。")

    add_heading(doc, "4.3 地圖與旅遊行程", 2)
    add_body(doc, "地圖除了顯示 BiteMap 已發布店家，也能另外呈現觀光署開放資料，包含餐飲、景點、旅館民宿及旅遊服務站。使用者可以從目前位置或指定區域查看附近地點，再依地區、價格與料理類型進行篩選。")
    add_body(doc, "找到適合的地點後，使用者可以直接加入我的行程：餐廳加入為吃飯地點，景點加入為景點，旅館民宿加入為住宿地點。加入後，可以在行程頁調整順序、刪除地點，並利用 BiteMap 內部地圖查看旅遊順序與位置連線。")
    add_body(doc, "因此，BiteMap 不只協助使用者發現地點，也能將搜尋結果直接轉換為實際可使用的旅遊行程。")

    add_heading(doc, "4.4 約飯 Meal", 2)
    add_body(doc, "約飯功能主要針對單人旅遊與陌生地區旅遊時的共同用餐需求。使用者可以從餐廳詳細頁直接建立約飯，並將正在查看的餐廳設定為候選餐廳。每一場約飯最多可以設定三家候選餐廳，再由參與者共同投票決定最後用餐地點。")
    add_bullet(doc, "公開約飯：符合條件的使用者可以直接加入。")
    add_bullet(doc, "私人約飯：使用者提出加入申請後，由發起人決定是否接受。")
    add_body(doc, "正式加入飯局後，成員可以進入約飯聊天室進行討論及即時溝通。約飯功能將找到餐廳與找到一起吃飯的人串在一起；對單人旅行者而言，不需要另外使用社群平台發文尋找旅伴，可以直接從餐廳選擇進入約飯流程。")

    add_heading(doc, "4.5 留言 再訪與社交功能", 2)
    add_body(doc, "BiteMap 不採用另一套傳統五星評分，而是使用會再訪、普通、不會三種狀態，讓使用者直接表達吃過之後是否願意再次前往。使用者再次造訪同一家店時，可以新增再訪紀錄；系統統計時只採用該使用者最新一次的狀態，因此同一個人不會因為多次造訪而重複計票。")
    add_body(doc, "除此之外，系統也提供收藏、留言、追蹤、好友及聊天等功能，使 BiteMap 不只是工具型地圖，也能逐漸累積屬於平台自己的在地美食資料與使用經驗。")

    add_heading(doc, "5 資料來源", 1)
    add_body(doc, "BiteMap 主要整合三種類型的資料，分別負責提供旅遊據點基礎資訊、大眾口碑與平台自身累積的實際使用經驗。")
    add_table(doc,
              ["資料來源", "主要內容", "在系統中的用途"],
              [
                  ("觀光署開放資料", "餐飲、景點、旅館民宿、旅遊服務站", "擴大旅遊據點涵蓋範圍，提供基礎且完整的地點資訊。"),
                  ("Google Places 資料", "Google 評分星數、Google 評論總數、Google Place ID", "提供外部大眾評價，作為熱門程度與大眾口碑的參考。"),
                  ("BiteMap 平台資料", "店家基本資料、使用者留言、再訪狀態、再訪紀錄、收藏、約飯與在地使用經驗", "補足官方資料與 Google 評分無法直接呈現的實際使用資訊。"),
              ],
              widths=[1.55, 3.0, 1.9])

    add_heading(doc, "5.1 觀光署開放資料", 2)
    add_body(doc, "觀光署資料主要提供全台旅遊據點的基礎資訊，目前系統規劃涵蓋餐飲資料、景點資料、旅館民宿資料及旅遊服務站資料。這些資料負責擴大 BiteMap 的旅遊地點涵蓋範圍，讓系統不只侷限於平台自行建立的餐廳資料。")

    add_heading(doc, "5.2 Google Places 資料", 2)
    add_body(doc, "Google Places API 主要提供外部大眾評價參考，目前 BiteMap 使用 Google 評分星數、Google 評論總數與 Google Place ID。Google 資料主要反映大量使用者對店家的整體評價，作為旅客判斷店家熱門程度與大眾口碑的參考。")

    add_heading(doc, "5.3 BiteMap 平台資料", 2)
    add_body(doc, "BiteMap 自有資料主要來自系統使用者與平台管理內容，包括店家基本資料、使用者留言、會再訪／普通／不會、再訪紀錄、收藏資料、約飯資料及使用者在地使用經驗。這些資料可以補足官方資料與 Google 評分無法直接呈現的資訊。")

    add_heading(doc, "資料整合特色", 2)
    add_body(doc, "BiteMap 的資料特色可以整理為：觀光署資料提供完整旅遊據點，Google 評價提供大眾口碑，BiteMap 再訪率與在地資料提供實際使用經驗。三種資料來源互相補充，讓使用者在搜尋餐廳、景點與住宿時，不只是看到哪裡有，也能進一步判斷值不值得去，以及其他人是否願意再去，進而完成旅遊決策。")

    add_heading(doc, "資料來源註記", 2)
    add_body(doc, "交通部觀光署政府資料開放平台資料集 7779、7777、7780、44055。", after=0)

    doc.core_properties.title = "BiteMap 系統目標與功能架構"
    doc.core_properties.subject = "美食探索 旅遊地圖 行程規劃與社交約飯服務"
    doc.core_properties.author = "BiteMap"
    doc.core_properties.comments = ""
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
